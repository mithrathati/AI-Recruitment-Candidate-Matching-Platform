import os
from io import BytesIO
from typing import Optional, Union
from app.config import settings
from app.utils.exceptions import (
    UnsupportedFileFormatError,
    EmptyDocumentError,
    CorruptedDocumentError,
    DocumentProcessingError,
)
from app.utils.file_utils import clean_text, file_extension


def _extract_pdf(source: Union[str, BytesIO, bytes]) -> str:
    try:
        import pdfplumber
    except ImportError as e:
        raise DocumentProcessingError("PDF processing library (pdfplumber) not installed") from e

    if isinstance(source, bytes):
        stream = BytesIO(source)
    elif isinstance(source, BytesIO):
        stream = source
        stream.seek(0)
    else:
        stream = source

    text_parts = []
    try:
        with pdfplumber.open(stream) as pdf:
            for page in pdf.pages:
                try:
                    t = page.extract_text() or ""
                except Exception:
                    t = ""
                text_parts.append(t)
    except Exception as e:
        if "password" in str(e).lower() or "encrypt" in str(e).lower():
            raise CorruptedDocumentError("PDF appears to be encrypted or corrupted and cannot be read.")
        raise CorruptedDocumentError(f"Failed to parse PDF file: {str(e)}")
    return "\n\n".join(text_parts)


def _extract_docx(source: Union[str, BytesIO, bytes]) -> str:
    try:
        from docx import Document
    except ImportError as e:
        raise DocumentProcessingError("DOCX processing library (python-docx) not installed") from e

    if isinstance(source, bytes):
        stream = BytesIO(source)
    elif isinstance(source, BytesIO):
        stream = source
        stream.seek(0)
    else:
        stream = source

    try:
        doc = Document(stream)
    except Exception as e:
        raise CorruptedDocumentError(f"Failed to parse DOCX file: {str(e)}")
    parts = []
    for p in doc.paragraphs:
        parts.append(p.text)
    for table in doc.tables:
        try:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells]
                parts.append(" | ".join(cells))
        except Exception:
            continue
    return "\n".join(parts)


def _extract_txt(source: Union[str, BytesIO, bytes]) -> str:
    if isinstance(source, bytes):
        try:
            return source.decode("utf-8", errors="replace")
        except Exception as e:
            raise CorruptedDocumentError(f"Failed to decode TXT content: {str(e)}")
    elif isinstance(source, BytesIO):
        try:
            source.seek(0)
            return source.read().decode("utf-8", errors="replace")
        except Exception as e:
            raise CorruptedDocumentError(f"Failed to read TXT stream: {str(e)}")
    else:
        try:
            with open(source, "r", encoding="utf-8", errors="replace") as f:
                return f.read()
        except Exception as e:
            raise CorruptedDocumentError(f"Failed to read TXT file: {str(e)}")


def extract_document_text(
    source: Union[str, bytes, BytesIO],
    filename_hint: Optional[str] = None,
) -> str:
    if isinstance(source, (bytes, BytesIO)):
        ext = file_extension(filename_hint or "")
        if not ext:
            raise UnsupportedFileFormatError(
                "Cannot determine file type from content alone. "
                "Please provide a filename hint or use a file with a known extension."
            )
    elif isinstance(source, str):
        if not os.path.exists(source):
            raise CorruptedDocumentError("Document file does not exist.")
        ext = file_extension(filename_hint or source)
    else:
        raise UnsupportedFileFormatError(f"Unsupported source type: {type(source).__name__}")
    raw = ""
    if ext == "pdf":
        raw = _extract_pdf(source)
    elif ext == "docx":
        raw = _extract_docx(source)
    elif ext == "txt":
        raw = _extract_txt(source)
    else:
        raise UnsupportedFileFormatError(f"Unsupported file extension: .{ext}")
    text = clean_text(raw)
    if not text or len(text) < 20:
        raise EmptyDocumentError(
            "The document appears empty or contains too little readable text. "
            "Please upload a valid resume/Job Description."
        )
    if len(text) < 50 and ext == "pdf":
        raise EmptyDocumentError(
            "The PDF may be a scanned image without extractable text. "
            "Please use a text-based PDF or DOCX."
        )
    return text


def validate_document(
    source: Union[str, bytes, BytesIO],
    filename_hint: Optional[str] = None,
) -> None:
    text = extract_document_text(source, filename_hint=filename_hint)
    if not text:
        raise EmptyDocumentError("Document has no readable text.")
    if len(text) > 1_000_000:
        raise DocumentProcessingError("Document is too large to process (>1M chars).")
